import os
import uuid
from werkzeug.utils import secure_filename
from flask import current_app


ALLOWED_EXTENSIONS = {'pdf', 'docx', 'doc', 'txt'}


def allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def save_uploaded_file(file) -> tuple:
    """Save uploaded file and return (saved_path, original_filename, file_type)."""
    original_filename = secure_filename(file.filename)
    ext = original_filename.rsplit('.', 1)[1].lower() if '.' in original_filename else 'bin'
    unique_filename = f"{uuid.uuid4().hex}.{ext}"
    upload_folder = current_app.config['UPLOAD_FOLDER']
    os.makedirs(upload_folder, exist_ok=True)
    save_path = os.path.join(upload_folder, unique_filename)
    file.save(save_path)
    return save_path, original_filename, ext


def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Extract raw text from PDF, DOCX, or TXT file."""
    try:
        if file_type == 'pdf':
            return _extract_from_pdf(file_path)
        elif file_type in ('docx', 'doc'):
            return _extract_from_docx(file_path)
        elif file_type == 'txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
    except Exception as e:
        current_app.logger.error(f"Error extracting text from {file_path}: {e}")
    return ''


def _extract_from_pdf(file_path: str) -> str:
    """Extract text from PDF using pdfplumber."""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return '\n'.join(text_parts)
    except Exception as e:
        # Fallback: try PyPDF2
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                return '\n'.join(
                    page.extract_text() or '' for page in reader.pages
                )
        except Exception:
            # Fallback 2: Sometimes files with .pdf extensions are just raw text files disguised as PDFs
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
                    if len(content.strip()) > 10:
                        return content
            except Exception:
                pass
            return ''


def _extract_from_docx(file_path: str) -> str:
    """Extract text from DOCX file."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return '\n'.join(paragraphs)
    except Exception:
        return ''


def delete_file(file_path: str) -> bool:
    """Delete a file safely."""
    try:
        if os.path.exists(file_path):
            os.remove(file_path)
            return True
    except Exception:
        pass
    return False
